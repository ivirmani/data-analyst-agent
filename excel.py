"""Collects tables while the agent works, then writes them into one Excel file.

Same shape as report.py, and for the same reason: every run_python call is a
fresh process, so tables are stashed in a JSON file on disk until the end.

The point of this over a CSV is native charts. What lands in the workbook is a
real Excel chart object, not a picture of one, so whoever opens it can click
into it, re-point it at other cells, and restyle it.
"""

import json
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DRAFT_FILE = ".excel_draft.json"

# Where the data starts. Row 1 holds the note, row 2 is blank, row 3 is the
# header, and the numbers begin on row 4.
NOTE_ROW = 1
HEADER_ROW = 3

# Excel refuses these characters in a sheet name, and caps the length at 31.
BAD_SHEET_CHARS = r"[]:*?/\\"

# Words that suggest a column holds a proportion rather than an amount.
SHARE_WORDS = ("share", "percent", "pct", "%", "rate", "proportion")


def clear():
    """Throw away any tables left over from a previous run."""
    if os.path.exists(DRAFT_FILE):
        os.remove(DRAFT_FILE)


def _load():
    if not os.path.exists(DRAFT_FILE):
        return []
    with open(DRAFT_FILE) as f:
        return json.load(f)


def add_table(name, df, note="", chart=None, labels=None, values=None):
    """Record one table for the workbook. Each becomes its own sheet.

    name:   the sheet name, for example "Revenue by region"
    df:     the DataFrame to write
    note:   a sentence describing the table, placed above it
    chart:  "bar" or "line" to add a real Excel chart, or leave it out
    labels: the column holding the category names, for example "Region"
    values: the column holding the numbers, for example "Total revenue"
    """
    tables = _load()
    tables.append(
        {
            "name": name,
            "note": note,
            "columns": [str(c) for c in df.columns],
            # tolist() turns numpy numbers into ordinary Python ones, which
            # json can write. Without it this raises "not JSON serializable".
            "rows": df.values.tolist(),
            "chart": chart,
            "labels": labels,
            "values": values,
        }
    )

    with open(DRAFT_FILE, "w") as f:
        json.dump(tables, f, indent=2, default=str)

    print(f"Recorded table {len(tables)}: {name}")


def _safe_sheet_name(name, used):
    """Make a sheet name Excel will accept, and that is not already taken."""
    for character in BAD_SHEET_CHARS:
        name = name.replace(character, "-")
    name = name[:31] or "Sheet"

    # If two tables want the same name, number them.
    candidate, counter = name, 2
    while candidate in used:
        suffix = f" {counter}"
        candidate = name[: 31 - len(suffix)] + suffix
        counter += 1
    return candidate


def _column_format(column_name, values):
    """Pick a display format for one column, from its name and its numbers."""
    numbers = [v for v in values if isinstance(v, (int, float))]
    if not numbers:
        return None

    name = str(column_name).lower()

    # A column called "share" holding values between -1 and 1 is a fraction.
    # Excel's % format multiplies by 100 on display, leaving the cell value
    # untouched, so formulas still see 0.504.
    if any(word in name for word in SHARE_WORDS):
        if all(-1 <= v <= 1 for v in numbers):
            return "0.0%"
        return "#,##0.0"

    if any(isinstance(v, float) for v in numbers):
        return "#,##0" if max(abs(v) for v in numbers) >= 1000 else "#,##0.00"
    return None


def _write_table(sheet, table):
    """Put one table onto one sheet, with a header and readable columns."""
    columns = table["columns"]

    sheet.cell(row=NOTE_ROW, column=1, value=table.get("note", "")).font = Font(
        italic=True, color="666666"
    )

    header_fill = PatternFill("solid", fgColor="EFEFEF")
    for index, column_name in enumerate(columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=column_name)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    # Work out one format per column, rather than per cell, so a column reads
    # consistently even when one value happens to be small.
    formats = [
        _column_format(name, [row[i] for row in table["rows"]])
        for i, name in enumerate(columns)
    ]

    for row_offset, row in enumerate(table["rows"], start=HEADER_ROW + 1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_offset, column=column_index, value=value)
            if formats[column_index - 1]:
                cell.number_format = formats[column_index - 1]

    # Widen each column to fit its longest entry, within reason.
    for index, column_name in enumerate(columns, start=1):
        longest = max(
            [len(str(column_name))] + [len(str(row[index - 1])) for row in table["rows"]]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 4, 40)

    sheet.cell(row=HEADER_ROW, column=1).alignment = Alignment(horizontal="left")


def _add_chart(sheet, table):
    """Attach a real Excel chart to the sheet, if the table asked for one."""
    kind = table.get("chart")
    if kind not in ("bar", "line"):
        return

    columns = table["columns"]
    label_name = table.get("labels")
    value_name = table.get("values")

    # If the columns named do not exist, skip the chart rather than crash.
    if label_name not in columns or value_name not in columns:
        print(f"  (skipped chart on '{table['name']}': columns not found)")
        return

    label_column = columns.index(label_name) + 1
    value_column = columns.index(value_name) + 1
    last_row = HEADER_ROW + len(table["rows"])

    chart = BarChart() if kind == "bar" else LineChart()
    chart.title = table["name"]
    chart.y_axis.title = value_name
    chart.x_axis.title = label_name
    chart.height = 8
    chart.width = 16

    # Reference points at a block of cells. titles_from_data uses the header
    # cell as the series name, which is why min_row starts on the header.
    data = Reference(sheet, min_col=value_column, min_row=HEADER_ROW, max_row=last_row)
    categories = Reference(
        sheet, min_col=label_column, min_row=HEADER_ROW + 1, max_row=last_row
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # Park it two columns to the right of the data.
    anchor = get_column_letter(len(columns) + 2) + str(HEADER_ROW)
    sheet.add_chart(chart, anchor)


def build(output="analysis.xlsx"):
    """Write every recorded table into one Excel workbook."""
    tables = _load()

    if not tables:
        print("No tables were recorded, so there is nothing to build.")
        return

    workbook = Workbook()
    workbook.remove(workbook.active)  # drop the empty default sheet

    used_names = set()
    for table in tables:
        name = _safe_sheet_name(table["name"], used_names)
        used_names.add(name)

        sheet = workbook.create_sheet(name)
        _write_table(sheet, table)
        _add_chart(sheet, table)

    workbook.save(output)
    clear()
    print(f"Workbook written to {output} ({len(tables)} sheets)")


# Build a small example workbook, so we can check this file works on its own.
if __name__ == "__main__":
    import pandas as pd

    clear()
    df = pd.read_csv("sales.csv")

    by_region = df.groupby("region")["revenue"].sum().round(2).reset_index()
    by_region.columns = ["Region", "Total revenue"]
    add_table(
        "Revenue by region",
        by_region,
        note="Total revenue for 2024, summed across every product and month.",
        chart="bar",
        labels="Region",
        values="Total revenue",
    )

    monthly = df.groupby("month")["revenue"].sum().round(2).reset_index()
    monthly.columns = ["Month", "Revenue"]
    add_table(
        "Monthly trend",
        monthly,
        note="Total revenue per month across all regions.",
        chart="line",
        labels="Month",
        values="Revenue",
    )

    build()
